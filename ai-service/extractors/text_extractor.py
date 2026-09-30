import os
import re
from io import BytesIO
from typing import List, Tuple

class TextExtractor:
    """Multi-format text extraction pipeline for PDF, DOCX, PPTX, and TXT files."""

    @staticmethod
    def extract_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, str]:
        ext = filename.split(".")[-1].lower() if "." in filename else "txt"
        raw_text = ""

        if ext == "txt":
            try:
                raw_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                raw_text = file_bytes.decode("latin-1", errors="ignore")

        elif ext == "pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(BytesIO(file_bytes))
                pages_text = []
                for idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    if text.strip():
                        pages_text.append(f"[Page {idx + 1}]\n{text.strip()}")
                raw_text = "\n\n".join(pages_text)
            except Exception as e:
                raw_text = f"[PDF Extraction Error: {str(e)}]\nFallback text representation for {filename}."

        elif ext == "docx":
            try:
                import docx
                doc = docx.Document(BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                raw_text = "\n\n".join(paragraphs)
            except Exception as e:
                raw_text = f"[DOCX Extraction Error: {str(e)}]\nFallback text representation for {filename}."

        elif ext == "pptx":
            try:
                from pptx import Presentation
                prs = Presentation(BytesIO(file_bytes))
                slides_text = []
                for idx, slide in enumerate(prs.slides):
                    slide_parts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_parts.append(shape.text.strip())
                    if slide_parts:
                        slides_text.append(f"[Slide {idx + 1}]\n" + "\n".join(slide_parts))
                raw_text = "\n\n".join(slides_text)
            except Exception as e:
                raw_text = f"[PPTX Extraction Error: {str(e)}]\nFallback text representation for {filename}."

        else:
            raw_text = file_bytes.decode("utf-8", errors="ignore")

        cleaned_text = TextExtractor.clean_and_normalize(raw_text)
        return raw_text, cleaned_text

    @staticmethod
    def clean_and_normalize(text: str) -> str:
        if not text:
            return ""
        # Remove consecutive blank lines and excess whitespace
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        return text.strip()

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1500, overlap: int = 200) -> List[str]:
        if not text:
            return []
        words = text.split()
        if len(words) <= chunk_size:
            return [text]

        chunks = []
        i = 0
        while i < len(words):
            chunk = " ".join(words[i:i + chunk_size])
            chunks.append(chunk)
            i += (chunk_size - overlap)
        return chunks
